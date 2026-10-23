"""
Step 13: Re-export error_analysis_cases as a wide-column Excel workbook with
all 5 system verdicts and reasoning, ready for manual error-category coding.

- Reads existing `data/results/error_analysis_cases.csv` (preserves the noteIds
  selected by step 9; doesn't re-sample).
- Joins fresh System D and System E columns from master_results.csv (the
  existing CSV may have been generated before D/E were added).
- Writes `data/results/error_analysis_cases.xlsx` with wide columns and
  text wrapping in the long-text cells.
"""

from pathlib import Path
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(r"C:\Users\yusuf\dissertation")
RESULTS = PROJECT_ROOT / "data" / "results"

# (column key, header, width, wrap)
COLUMNS = [
    ("noteId",                     "Note ID",                  18, False),
    ("tweetId",                    "Tweet ID",                 18, False),
    ("category",                   "Category",                 12, False),
    ("cn_verdict",                 "CN baseline",              16, False),
    ("extracted_claim",            "Extracted claim",          50, True),
    ("summary",                    "Original CN summary",      60, True),
    ("gpt4o_verdict",              "A: GPT-4o verdict",        14, False),
    ("gpt4o_confidence",           "A: confidence",            12, False),
    ("gpt4o_reasoning",            "A: reasoning",             65, True),
    ("tavily_gpt4o_verdict",       "B: Tavily RAG verdict",    14, False),
    ("tavily_gpt4o_confidence",    "B: confidence",            12, False),
    ("tavily_gpt4o_reasoning",     "B: reasoning",             65, True),
    ("google_verdict",             "C: Google FC verdict",     14, False),
    ("google_source",              "C: source",                20, False),
    ("google_raw_rating",          "C: raw rating",            22, True),
    ("filtered_rag_verdict",       "D: Filtered RAG verdict",  14, False),
    ("filtered_rag_confidence",    "D: confidence",            12, False),
    ("filtered_rag_reasoning",     "D: reasoning",             65, True),
    ("llama_verdict",              "E: Llama verdict",         14, False),
    ("llama_confidence",           "E: confidence",            12, False),
    ("llama_reasoning",            "E: reasoning",             65, True),
    ("error_category",             "ERROR CATEGORY (fill in)", 28, False),
    ("notes",                      "NOTES (fill in)",          50, True),
]


def main() -> int:
    cases_path = RESULTS / "error_analysis_cases.csv"
    master_path = RESULTS / "master_results.csv"
    if not cases_path.exists():
        print(f"ERROR: {cases_path} not found. Run 09_error_analysis.py first.", file=sys.stderr)
        return 1

    cases = pd.read_csv(cases_path, dtype={"noteId": str, "tweetId": str})
    master = pd.read_csv(master_path, dtype={"noteId": str})
    print(f"  loaded {len(cases)} error-analysis cases", flush=True)

    # Pull the System D + E columns from master if they're missing in cases.
    needed = ["filtered_rag_verdict", "filtered_rag_confidence", "filtered_rag_reasoning",
              "llama_verdict", "llama_confidence", "llama_reasoning"]
    extra = [c for c in needed if c not in cases.columns]
    if extra:
        keep = ["noteId"] + [c for c in needed if c in master.columns]
        cases = cases.merge(master[keep], on="noteId", how="left")
        print(f"  joined missing columns from master_results: {extra}", flush=True)

    if "error_category" not in cases.columns:
        cases["error_category"] = ""
    if "notes" not in cases.columns:
        cases["notes"] = ""

    # Persist the merged form back as CSV too so downstream stays consistent
    cases.to_csv(cases_path, index=False)

    # ---- Build the Excel workbook ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Error Analysis"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    wrap_align  = Alignment(wrap_text=True, vertical="top")
    plain_align = Alignment(vertical="top")

    # Headers
    for ci, (key, header, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=ci, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "F2"  # freeze first 5 columns + header so claim/reasoning scrolls

    # Body
    for ri, (_, row) in enumerate(cases.iterrows(), start=2):
        for ci, (key, _, _, wrap) in enumerate(COLUMNS, start=1):
            val = row.get(key, "")
            if pd.isna(val): val = ""
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.alignment = wrap_align if wrap else plain_align
        ws.row_dimensions[ri].height = 110  # tall rows for readable text wrapping

    out_path = RESULTS / "error_analysis_cases.xlsx"
    wb.save(out_path)
    print(f"\n  saved {out_path}", flush=True)
    print(f"  rows: {len(cases)}, columns: {len(COLUMNS)}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
